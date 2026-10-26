"""Small synthetic archives reproduce expansion and dimension attacks safely."""
import io
import zipfile
from types import SimpleNamespace

import openpyxl
import pytest
from modules.finance import statement_import as parser


def test_excel_expansion_limit_precedes_workbook_decoder(monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/sharedStrings.xml", b"x" * 4096)
    monkeypatch.setattr(parser, "MAX_XLSX_EXPANDED_BYTES", 2048)
    def forbidden(*a, **k):
        pytest.fail("Workbook decoder ran before expansion checks")
    monkeypatch.setattr(openpyxl, "load_workbook", forbidden)
    with pytest.raises(parser.StatementError, match="expanded") as error:
        parser.parse_statement("small.xlsx", buffer.getvalue())
    assert error.value.status == 413


def test_excel_member_limit_precedes_workbook_decoder(monkeypatch):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for n in range(3):
            archive.writestr(f"part-{n}.xml", b"x")
    monkeypatch.setattr(parser, "MAX_XLSX_MEMBERS", 2)
    with pytest.raises(parser.StatementError) as error:
        parser.parse_statement("parts.xlsx", buffer.getvalue())
    assert error.value.status == 413


def test_excel_sparse_extreme_dimensions_are_rejected():
    workbook = openpyxl.Workbook()
    workbook.active.append(["Date", "Description", "Amount"])
    workbook.active.cell(row=1, column=16384, value="far-away formatted cell")
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    with pytest.raises(parser.StatementError, match="too large") as error:
        parser.parse_statement("sparse.xlsx", buffer.getvalue())
    assert error.value.status == 413


def test_excel_workbook_is_closed_when_cell_budget_rejects(monkeypatch):
    workbook = openpyxl.Workbook()
    workbook.active.append(["Date", "Description", "Amount"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    closed = []
    original = openpyxl.load_workbook
    def load(*args, **kwargs):
        actual = original(*args, **kwargs)
        close = actual.close
        actual.close = lambda: (closed.append(True), close())
        return actual
    monkeypatch.setattr(openpyxl, "load_workbook", load)
    monkeypatch.setattr(parser, "MAX_GRID_CELLS", 2)
    with pytest.raises(parser.StatementError):
        parser.parse_statement("cells.xlsx", buffer.getvalue())
    assert closed == [True]


def test_legacy_excel_dimensions_checked_and_resources_released(monkeypatch):
    import xlrd
    released = []
    sheet = SimpleNamespace(nrows=2, ncols=10000)
    book = SimpleNamespace(nsheets=1, sheet_by_index=lambda i: sheet,
                           release_resources=lambda: released.append(True))
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: book)
    with pytest.raises(parser.StatementError) as error:
        parser.parse_statement("old.xls", b"synthetic")
    assert error.value.status == 413 and released == [True]


def test_legacy_excel_dates_and_sheet_unloading(monkeypatch):
    import xlrd
    unloaded, released = [], []
    rows = [[SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value=v) for v in ("Date", "Narration", "Amount")],
            [SimpleNamespace(ctype=xlrd.XL_CELL_DATE, value=45383),
             SimpleNamespace(ctype=xlrd.XL_CELL_TEXT, value="SWIGGY"),
             SimpleNamespace(ctype=xlrd.XL_CELL_NUMBER, value=50)]]
    sheet = SimpleNamespace(nrows=2, ncols=3, row=lambda i: rows[i])
    book = SimpleNamespace(nsheets=1, datemode=0, sheet_by_index=lambda i: sheet,
                           unload_sheet=unloaded.append, release_resources=lambda: released.append(True))
    monkeypatch.setattr(xlrd, "open_workbook", lambda **kwargs: book)
    result = parser.parse_statement("old.xls", b"synthetic")
    assert result["rows"][0]["date"] == "2024-04-01"
    assert unloaded == [0] and released == [True]


@pytest.mark.parametrize("limit", ["pages", "stream", "text"])
def test_pdf_limits_reject_before_downstream_work(monkeypatch, limit):
    import pypdf
    called = []
    def extract(**kwargs):
        called.append(True)
        return "x" * 30
    page = SimpleNamespace(get_contents=lambda: SimpleNamespace(get_data=lambda: b"1234"), extract_text=extract)
    reader = SimpleNamespace(is_encrypted=False, pages=[page, page])
    monkeypatch.setattr(pypdf, "PdfReader", lambda *a, **k: reader)
    monkeypatch.setattr(parser, {"pages": "MAX_PDF_PAGES", "stream": "MAX_PDF_PAGE_STREAM_BYTES", "text": "MAX_PDF_TEXT_CHARS"}[limit], 1)
    with pytest.raises(parser.StatementError) as error:
        parser.parse_statement("document.pdf", b"synthetic")
    assert error.value.status == 413
    assert len(called) == (1 if limit == "text" else 0)


@pytest.mark.parametrize("filename,data", [("wide.csv", b"a,b,c,d\n1,2,3,4"),
                                            ("wide.json", b'[{"date":1,"amount":2,"a":3,"b":4}]')])
def test_text_grid_bounds_precede_matrix_allocation(monkeypatch, filename, data):
    monkeypatch.setattr(parser, "MAX_GRID_COLUMNS", 3)
    with pytest.raises(parser.StatementError) as error:
        parser.parse_statement(filename, data)
    assert error.value.status == 413
