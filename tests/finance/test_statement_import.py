"""Statement import: pure parsers + the /parse -> /import API round trip. No real mic, no network."""
import io
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from api.main import app
from modules.finance import statement_import as si
from shared import user_store

client = TestClient(app)
USER = "__pytest_statement_import__"

HDFC = (
    "HDFC BANK Ltd. Statement of account\n"
    "Account No: XXXX1234\n\n"
    "Date,Narration,Chq./Ref.No.,Value Dt,Withdrawal Amt.,Deposit Amt.,Closing Balance\n"
    "01/04/24,UPI-SWIGGY-swiggy@icici-ICIC0001-412345678901-Payment,0000412345678901,01/04/24,450.00,,49550.00\n"
    "01/04/24,Opening Balance,,,,,50000.00\n"
    "05/04/24,NEFT CR-ACME CORP-SALARY APR,N123,05/04/24,,\"80,000.00\",129550.00\n"
    "10/04/24,ACH D- HDFC LOAN-00123,A1,10/04/24,\"12,000.00\",,117550.00\n"
    "12/04/24,POS ZERODHA BROKING SIP,B1,12/04/24,5000.00,,112550.00\n"
    "20/04/24,RENT TO LANDLORD,C1,20/04/24,\"20,000.00\",,92550.00\n"
    "28/04/24,UBER INDIA,D1,28/04/24,300.00,,92250.00\n"
    "30/04/24,Total,,,\"37,750.00\",\"80,000.00\",\n"
    "Date,Narration,Chq./Ref.No.,Value Dt,Withdrawal Amt.,Deposit Amt.,Closing Balance\n"
)


@pytest.fixture(autouse=True)
def _clean():
    user_store.clear_transactions(USER)
    user_store.delete_financial_profile(USER)
    yield
    user_store.clear_transactions(USER)
    user_store.delete_financial_profile(USER)


def test_hdfc_csv_preamble_columns_skips_and_categories():
    res = si.parse_statement("stmt.csv", HDFC.encode())
    rows = res["rows"]
    assert len(rows) == 6 and res["skipped"] >= 3  # opening balance, total, repeated header
    by = {r["merchant"]: r for r in rows}
    assert by["Swiggy"]["category"] == "food" and by["Swiggy"]["amount"] == 450 and by["Swiggy"]["date"] == "2024-04-01"
    assert by["Swiggy"]["direction"] == "debit"
    assert [r["category"] for r in rows if r["direction"] == "credit"] == ["income"]
    cats = {r["category"] for r in rows}
    assert {"emi", "investment", "rent", "transport"} <= cats


def test_signed_amount_csv_and_tsv_and_semicolon():
    csv_text = "Txn Date,Description,Amount\n2024-05-01,Zomato order,-250.50\n2024-05-02,Payroll,\"50,000\"\n"
    rows = si.parse_statement("a.csv", csv_text.encode())["rows"]
    assert [(r["direction"], r["amount"]) for r in rows] == [("debit", 250.5), ("credit", 50000.0)]
    tsv = csv_text.replace(",", "\t").replace('"', "")
    assert len(si.parse_statement("a.tsv", tsv.replace("50\t000", "50000").encode())["rows"]) == 2
    semi = "Txn Date;Description;Amount\n2024-05-01;Zomato order;-250.50\n"
    assert len(si.parse_statement("a.csv", semi.encode())["rows"]) == 1


def test_drcr_marker_column_and_latin1():
    txt = "Date,Particulars,Amount,Dr/Cr\n01-Jan-2024,Caf\xe9 Coffee,100,Dr\n02-Jan-2024,Refund,40,Cr\n"
    rows = si.parse_statement("m.csv", txt.encode("latin-1"))["rows"]
    assert [r["direction"] for r in rows] == ["debit", "credit"]


def test_xlsx_with_excel_serial_dates_and_second_sheet():
    wb = openpyxl.Workbook()
    wb.active.append(["nothing useful here"])
    ws = wb.create_sheet("Txns")
    ws.append(["Statement for April"])
    ws.append(["Date", "Narration", "Debit", "Credit"])
    ws.append([45383, "BIGBASKET GROCERIES", 1200, None])  # 2024-04-01
    ws.append([45384, "SALARY", None, 60000])
    buf = io.BytesIO()
    wb.save(buf)
    res = si.parse_statement("s.xlsx", buf.getvalue())
    assert res["format"] == "xlsx"
    assert [(r["date"], r["category"]) for r in res["rows"]] == [("2024-04-01", "food"), ("2024-04-02", "income")]


def _pdf(lines: list[str]) -> bytes:
    objs = []
    stream = "BT /F1 10 Tf 40 780 Td 12 TL " + " ".join(f"({ln}) Tj T*" for ln in lines) + " ET"
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append("<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
    objs.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 600 800] /Contents 4 0 R "
                "/Resources << /Font << /F1 5 0 R >> >> >>")
    objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    return out + f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


def test_text_pdf_lines_use_balance_delta_for_direction():
    pdf = _pdf(["Statement of account", "01/04/2024 UPI SWIGGY 450.00 9550.00", "05/04/2024 SALARY ACME 5000.00 14550.00",
                "07/04/2024 Closing Balance 14550.00"])
    res = si.parse_statement("s.pdf", pdf)
    assert [(r["direction"], r["amount"]) for r in res["rows"]] == [("debit", 450.0), ("credit", 5000.0)]


def test_pdf_free_text_parser_dr_cr_suffix():
    raws, _ = si._text_rows(["12-Apr-24 Netflix 649.00 Dr", "garbage line", "13 Apr 2024 REFUND 10.00 Cr"])
    assert [(r["dir"], r["amt"]) for r in raws] == [("debit", 649.0), ("credit", 10.0)]


def test_image_only_pdf_and_garbage_give_friendly_errors():
    with pytest.raises(si.StatementError, match="scan or image"):
        si.parse_statement("s.pdf", _pdf([]))
    with pytest.raises(si.StatementError, match="could not be read"):
        si.parse_statement("s.pdf", b"not a pdf")
    with pytest.raises(si.StatementError, match="couldn't find any transactions"):
        si.parse_statement("s.txt", b"hello world\nnothing here")


def test_json_and_ofx():
    js = json.dumps([{"date": "2024-03-01", "description": "Spotify", "amount": -119}])
    assert si.parse_statement("t.json", js.encode())["rows"][0]["category"] == "entertainment"
    ofx = "<STMTTRN><DTPOSTED>20240301<TRNAMT>-500.00<NAME>PVR CINEMAS</STMTTRN>"
    assert si.parse_statement("t.ofx", ofx.encode())["rows"][0]["merchant"] == "Pvr Cinemas"


def test_xls_without_engine_message(monkeypatch):
    with pytest.raises(si.StatementError, match="xlsx or CSV"):
        si.parse_statement("old.xls", b"\xd0\xcf\x11\xe0 junk")


def test_header_detection_and_edge_parsing():
    grid = [["x"]] * 5 + [["Transaction Date", "Transaction Remarks", "Withdrawal", "Deposit"]]
    assert si.find_header(grid) == (5, {"date": 0, "desc": 1, "debit": 2, "credit": 3})
    assert si.find_header([["a", "b"]]) is None
    assert si.parse_amount("₹ 1,234.50") == 1234.5 and si.parse_amount("(500.00)") == -500
    assert si.parse_amount("750 Dr") == -750 and si.parse_amount("750 Cr") == 750 and si.parse_amount("n/a") is None
    assert si.parse_date("03/04/2024") == "2024-04-03"  # day-first
    assert si.parse_date("3-Apr-24") == "2024-04-03" and si.parse_date("45383") is None
    assert si.parse_date(45383) == "2024-04-01" and si.parse_date("garbage") is None
    assert si.clean_merchant("UPI/412345678901/ZOMATO LTD/zomato@hdfc/HDFC0000001") == "Zomato Ltd"
    assert len(si.clean_merchant("A" * 100 + " " + "B" * 50)) <= 40


def test_categorise_rules():
    c = lambda n, d="debit": si.categorize(n, d)  # noqa: E731
    assert (c("Airtel broadband"), c("Vi prepaid"), c("Apollo Pharmacy"), c("Udemy course")) == (
        "utilities", "utilities", "health", "education")
    assert (c("Hotstar"), c("Random shop"), c("Salary", "credit"), c("Friend", "credit")) == (
        "entertainment", "other", "income", "credit")
    assert c("Olive garden") == "other"  # \bola\b must not match 'olive'


def test_dedupe_within_file_and_unsigned_amounts_warning():
    text = "Date,Description,Amount\n01/02/2024,Coffee,100\n01/02/2024,Coffee,100\n02/02/2024,Salary,5000\n"
    res = si.parse_statement("d.csv", text.encode())
    assert len(res["rows"]) == 2 and res["skipped"] == 1 and res["warnings"]
    assert [r["direction"] for r in res["rows"]] == ["debit", "credit"]


def test_summary_monthly_averages_and_suggested_profile():
    rows = [
        {"date": "2024-04-02", "merchant": "A", "amount": 1000.0, "direction": "debit", "category": "food"},
        {"date": "2024-05-02", "merchant": "B", "amount": 3000.0, "direction": "debit", "category": "food"},
        {"date": "2024-05-03", "merchant": "C", "amount": 500.0, "direction": "debit", "category": "emi"},
        {"date": "2024-05-04", "merchant": "D", "amount": 800.0, "direction": "debit", "category": "investment"},
        {"date": "2024-04-05", "merchant": "S", "amount": 50000.0, "direction": "credit", "category": "income"},
        {"date": "2024-05-05", "merchant": "S", "amount": 50000.0, "direction": "credit", "category": "income"},
        {"date": "2024-05-06", "merchant": "G", "amount": 100.0, "direction": "credit", "category": "credit"},
    ]
    s = si.summarize(rows)
    assert s["period"] == {"from": "2024-04-02", "to": "2024-05-06", "months": 2}
    assert s["summary"]["expenses_by_category"] == {"food": 4000.0}
    assert s["summary"]["monthly_avg"] == {"food": 2000.0}
    assert s["summary"]["total_income"] == 100100.0 and s["summary"]["avg_monthly_income"] == 50050.0
    assert s["summary"]["emi_total"] == 500.0 and s["summary"]["investment_total"] == 800.0
    assert s["suggested_profile"] == {"expenses": {"food": 2000}, "monthly_income": 50000}
    short = si.summarize(rows[:1])
    assert short["suggested_profile"]["expenses"] == {}  # under a month of data


def _post(name: str, data: bytes):
    return client.post("/api/transactions/parse", files={"file": (name, data)}, data={"session_id": USER})


def test_api_unsupported_oversize_and_bad_file():
    r = _post("a.exe", b"x")
    assert r.status_code == 415 and ".csv" in r.json()["detail"]
    assert _post("a.csv", b"x" * (5 * 1024 * 1024 + 1)).status_code == 413
    r = _post("a.csv", b"nothing")
    assert r.status_code == 422 and "Traceback" not in r.text


def test_api_round_trip_parse_then_import():
    r = _post("hdfc.csv", HDFC.encode())
    assert r.status_code == 200
    body = r.json()
    assert body["rows_total"] == 6 and body["format"] == "csv" and body["period"]["months"] == 1
    assert body["suggested_profile"]["expenses"] == {"food": 450, "rent": 20000, "transport": 300}
    assert body["suggested_profile"]["monthly_income"] == 80000
    user_store.add_transaction(USER, date="2024-04-28", category="transport", amount=300.0, merchant="Uber India")
    imp = client.post("/api/transactions/import", json={"session_id": USER, "import_id": body["import_id"]})
    assert imp.json() == {"imported": 4, "duplicates": 1, "profile_updated": False}  # swiggy, rent, emi, salary (+1 dup)
    stored = {t["category"] for t in user_store.get_transactions(USER)}
    assert stored == {"food", "housing", "emi", "transport"}
    assert user_store.get_financial_profile(USER) is None
    imp = client.post("/api/transactions/import",
                      json={"session_id": USER, "import_id": body["import_id"], "apply_to_profile": True})
    assert imp.json() == {"imported": 0, "duplicates": 5, "profile_updated": True}
    prof = user_store.get_financial_profile(USER)
    assert prof["monthly_income"] == 80000 and prof["expenses"]["rent"] == 20000
    assert prof["sources"]["expenses"]["src"] == "form"


def test_import_rejects_unknown_or_other_session():
    body = _post("hdfc.csv", HDFC.encode()).json()
    assert client.post("/api/transactions/import", json={"session_id": "other", "import_id": body["import_id"]}).status_code == 404
    assert client.post("/api/transactions/import", json={"session_id": USER, "import_id": "nope"}).status_code == 404


_LAYOUT = [
    "Date        Narration                 Withdrawal   Deposit     Balance",
    "01/06/2026  NEFT CR-ACME CORP-SALARY  0.00         120,000.00  320,000.00",
    "03/06/2026  UPI-LANDLORD RENT-9999    30,000.00    0.00        290,000.00",
    "                                      (June)",
    "04/06/2026  ACH D- HDFC LTD EMI       15,000.00    0.00        275,000.00",
]


def _text_pdf(lines: list[str]) -> bytes:
    """Minimal one-page PDF in monospaced Courier (no reportlab): pypdf layout mode keeps the columns."""
    body = "BT /F1 9 Tf 12 TL 20 800 Td " + " T* ".join(f"({s}) Tj" for s in lines) + " ET"
    objs = ["<</Type/Catalog/Pages 2 0 R>>", "<</Type/Pages/Kids[3 0 R]/Count 1>>",
            "<</Type/Page/Parent 2 0 R/MediaBox[0 0 700 842]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>",
            f"<</Length {len(body)}>>\nstream\n{body}\nendstream", "<</Type/Font/Subtype/Type1/BaseFont/Courier>>"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode() + b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    return out + f"trailer<</Size {len(objs) + 1}/Root 1 0 R>>\nstartxref\n{xref}\n%%EOF".encode()


def test_text_pdf_layout_columns_decide_debit_credit():
    r = si.parse_statement("s.pdf", _text_pdf(_LAYOUT))
    got = [(x["direction"], x["amount"]) for x in r["rows"]]
    assert got == [("credit", 120000.0), ("debit", 30000.0), ("debit", 15000.0)]
    assert r["rows"][0]["category"] == "income" and r["rows"][2]["category"] == "emi"


def test_text_rows_glued_amounts_and_wrapped_narration():
    rows, _ = si._text_rows([
        "01/06/2026  SOMETHING ODD  0.00  120,000.00320,000.00",
        "03/06/2026  UPI-LANDLORD   30,000.000.00 290,000.00",
        "            RENT JUNE",
        "04/06/2026  SHOP  500.00Dr",
    ])
    assert [(r["dir"], r["amt"]) for r in rows] == [("credit", 120000.0), ("debit", 30000.0), ("debit", 500.0)]
    assert rows[1]["narr"].endswith("RENT JUNE")


def test_applying_a_statement_to_the_profile_also_records_its_monthly_emi(monkeypatch):
    """Spending averages exclude EMIs, so the profile must get the EMI too or the surplus is overstated."""
    from fastapi.testclient import TestClient

    from api.main import app
    from shared import user_store

    client, owner = TestClient(app), "emi-apply-owner"
    csv = "Date,Narration,Debit,Credit\n" + "\n".join(
        f"{d:02d}/{m:02d}/2026,{n},{a},{c}" for m in (6, 7) for d, n, a, c in (
            (1, "SALARY ACME", "", 100000), (3, "ACH D HDFC LTD EMI", 12000, ""), (5, "SWIGGY", 800, ""), (9, "BIGBASKET", 1500, ""),
            (12, "BESCOM", 900, ""), (15, "UBER", 400, ""), (18, "NETFLIX", 649, ""), (20, "RENT LANDLORD", 20000, ""), (22, "APOLLO PHARMACY", 700, ""),
            (25, "AMAZON", 1200, ""), (27, "ZOMATO", 500, "")))
    user_store.delete_financial_profile(owner)
    r = client.post("/api/transactions/parse", data={"session_id": owner}, files={"file": ("s.csv", csv.encode(), "text/csv")}).json()
    out = client.post("/api/transactions/import", json={"session_id": owner, "import_id": r["import_id"], "apply_to_profile": True}).json()
    assert out["profile_updated"]
    prof = user_store.get_financial_profile(owner)
    assert prof["loans"] and prof["loans"][0]["emi"] == 12000


def test_pdf_date_glued_to_narration():
    raws, _ = si._text_rows(["01/06/2026NEFT CR-ACME CORP-SALARY 80,000.00 1,80,000.00", "02/06/2026UPI SWIGGY 450.00 1,79,550.00"])
    assert [(r["date"], r["dir"], r["amt"]) for r in raws] == [("2026-06-01", "credit", 80000.0), ("2026-06-02", "debit", 450.0)]


def test_real_xls_file():
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("S")
    for i, row in enumerate([["Date", "Narration", "Withdrawal Amt.", "Deposit Amt."], ["01/04/2024", "UPI SWIGGY", 450, ""], ["05/04/2024", "NEFT SALARY ACME", "", 80000]]):
        for j, v in enumerate(row):
            ws.write(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    res = si.parse_statement("old.xls", buf.getvalue())
    assert [(r["category"], r["amount"]) for r in res["rows"]] == [("food", 450.0), ("income", 80000.0)]


def test_import_stores_salary_credits_as_income_and_skips_duplicates():
    csv_ = b"Date,Narration,Withdrawal Amt.,Deposit Amt.\n" + b"".join(
        f"0{m}/0{m}/2026,UPI SHOP{m},{100 * m}.00,\n2{m}/0{m}/2026,NEFT SALARY ACME,,50000.00\n".encode() for m in range(1, 4))
    for _ in range(2):
        p = client.post("/api/transactions/parse", files={"file": ("s.csv", csv_)}, data={"session_id": USER}).json()
        out = client.post("/api/transactions/import", json={"session_id": USER, "import_id": p["import_id"]}).json()
    assert out == {"imported": 0, "duplicates": 6, "profile_updated": False}
    assert len(user_store.get_transactions(USER)) == 3 and len(user_store.get_transactions(USER, kind="income")) == 3
    assert len(user_store.get_transactions(USER, kind=None)) == 6
